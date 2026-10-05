"""
tests/test_file_parser.py
~~~~~~~~~~~~~~~~~~~~~~~~~
Covers the upload-side text extraction behind Resume Studio.

Everything downstream — the structured document, the LaTeX, the PDF — is built
from this text alone, so anything the extractor drops is gone for good. These
tests pin the three things a plain "join the paragraphs" pass loses: table
cells, hyperlink targets and list markers.
"""

from __future__ import annotations

import io

import pytest

from api.utils.file_parser import extract_text_from_bytes

docx = pytest.importorskip("docx")
pymupdf = pytest.importorskip("pymupdf")


def _add_hyperlink(paragraph: object, text: str, url: str) -> None:
    """python-docx has no public API for writing a hyperlink."""
    from docx.opc.constants import RELATIONSHIP_TYPE
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    r_id = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)  # type: ignore[attr-defined]
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    link.append(run)
    paragraph._p.append(link)  # type: ignore[attr-defined]


def _docx_bytes(document: object) -> bytes:
    buffer = io.BytesIO()
    document.save(buffer)  # type: ignore[attr-defined]
    return buffer.getvalue()


class TestDocx:
    def test_hyperlink_targets_survive(self) -> None:
        document = docx.Document()
        paragraph = document.add_paragraph("+91-9000000000 | ")
        _add_hyperlink(paragraph, "LinkedIn", "https://www.linkedin.com/in/ada")

        text = extract_text_from_bytes(_docx_bytes(document), "resume.docx")

        assert "LinkedIn" in text
        assert "https://www.linkedin.com/in/ada" in text

    def test_a_link_whose_text_is_already_the_url_is_not_repeated(self) -> None:
        document = docx.Document()
        paragraph = document.add_paragraph("")
        _add_hyperlink(paragraph, "github.com/ada/rag", "https://www.github.com/ada/rag")

        text = extract_text_from_bytes(_docx_bytes(document), "resume.docx")

        assert text.count("github.com/ada/rag") == 1

    def test_table_cells_are_read_in_document_order(self) -> None:
        document = docx.Document()
        document.add_paragraph("EXPERIENCE")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Senior Engineer"
        table.cell(0, 1).text = "2022 - Present"
        document.add_paragraph("EDUCATION")

        text = extract_text_from_bytes(_docx_bytes(document), "resume.docx")

        assert "Senior Engineer" in text
        assert "2022 - Present" in text
        assert text.index("EXPERIENCE") < text.index("Senior Engineer") < text.index("EDUCATION")

    def test_list_paragraphs_keep_a_bullet_marker(self) -> None:
        document = docx.Document()
        document.add_paragraph("Stripe - Senior Engineer")
        document.add_paragraph("Cut p99 latency 38%", style="List Bullet")

        lines = extract_text_from_bytes(_docx_bytes(document), "resume.docx").splitlines()

        assert "Stripe - Senior Engineer" in lines
        assert "- Cut p99 latency 38%" in lines


class TestPdf:
    def test_link_targets_survive(self) -> None:
        document = pymupdf.open()
        page = document.new_page()
        page.insert_text((72, 72), "Ada Lovelace")
        page.insert_text((72, 100), "GitHub")
        page.insert_link({
            "kind": pymupdf.LINK_URI,
            "from": pymupdf.Rect(70, 88, 120, 104),
            "uri": "https://github.com/ada",
        })

        text = extract_text_from_bytes(document.tobytes(), "resume.pdf")

        assert "Ada Lovelace" in text
        assert "GitHub: https://github.com/ada" in text
